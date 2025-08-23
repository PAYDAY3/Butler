import pandas as pd
import numpy as np
import os
import glob
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, Reference
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import StandardScaler
from scipy import stats
from sklearn.cluster import KMeans
from sklearn.decomposition import PCA

def load_excel_files(directory):
    """加载目录中的所有Excel文件"""
    all_data = []
    excel_files = glob.glob(os.path.join(directory, "*.xlsx"))
    for file in excel_files:
        df = pd.read_excel(file, sheet_name=None)
        for sheet_name, sheet_data in df.items():
            sheet_data['源文件'] = os.path.basename(file)
            sheet_data['工作表'] = sheet_name
            all_data.append(sheet_data)
    return pd.concat(all_data, ignore_index=True)

def clean_data(df):
    """清理数据：删除重复行，处理缺失值，移除异常值"""
    # 删除重复行
    df.drop_duplicates(inplace=True)
    
    # 处理缺失值
    numeric_columns = df.select_dtypes(include=['int64', 'float64']).columns
    categorical_columns = df.select_dtypes(include=['object']).columns
    
    # 用中位数填充缺失的数值
    imputer = SimpleImputer(strategy='median')
    df[numeric_columns] = imputer.fit_transform(df[numeric_columns])
    
    # 用众数填充缺失的分类值
    df[categorical_columns] = df[categorical_columns].fillna(df[categorical_columns].mode().iloc[0])
    
    # 移除异常值（使用Z-分数方法）
    for column in numeric_columns:
        z_scores = np.abs(stats.zscore(df[column]))
        df = df[(z_scores < 3)]
    
    return df

def filter_data(df, column, condition):
    """根据给定的条件筛选数据"""
    if condition.startswith(">"):
        return df[df[column] > float(condition[1:])]
    elif condition.startswith("<"):
        return df[df[column] < float(condition[1:])]
    elif condition.startswith("="):
        return df[df[column] == condition[1:]]
    else:
        return df[df[column].str.contains(condition, na=False)]

def merge_data(df, merge_column):
    """根据指定列合并数据"""
    return df.groupby(merge_column).agg({col: 'sum' if df[col].dtype in ['int64', 'float64'] else 'first' 
                                         for col in df.columns if col != merge_column}).reset_index()

def analyze_data(df):
    """对数据进行分析"""
    numeric_columns = df.select_dtypes(include=['int64', 'float64']).columns
    analysis = {}
    
    # 基本统计
    analysis['基本统计'] = df[numeric_columns].describe()
    
    # 相关性矩阵
    analysis['相关性矩阵'] = df[numeric_columns].corr()
    
    # 偏度和峰度
    analysis['偏度'] = df[numeric_columns].skew()
    analysis['峰度'] = df[numeric_columns].kurtosis()
    
    # 分类变量的唯一值计数
    categorical_columns = df.select_dtypes(include=['object']).columns
    analysis['分类变量计数'] = {col: df[col].value_counts() for col in categorical_columns}
    
    # 执行K-means聚类
    if len(numeric_columns) >= 2:
        kmeans = KMeans(n_clusters=3, random_state=42)
        df['聚类'] = kmeans.fit_predict(df[numeric_columns])
        analysis['K-means聚类'] = df['聚类'].value_counts()
    
    # 执行PCA
    if len(numeric_columns) >= 2:
        pca = PCA(n_components=2)
        pca_result = pca.fit_transform(df[numeric_columns])
        df['PCA1'] = pca_result[:, 0]
        df['PCA2'] = pca_result[:, 1]
        analysis['PCA解释方差比'] = pca.explained_variance_ratio_
    
    return analysis, df

def create_visualizations(df, analysis, output_dir):
    """创建可视化图表"""
    numeric_columns = df.select_dtypes(include=['int64', 'float64']).columns
    categorical_columns = df.select_dtypes(include=['object']).columns
    
    # 数值列的直方图
    for column in numeric_columns:
        plt.figure(figsize=(10, 6))
        sns.histplot(df[column], kde=True)
        plt.title(f"{column}的分布")
        plt.savefig(os.path.join(output_dir, f"{column}_直方图.png"))
        plt.close()
    
    # 数值列的箱线图
    plt.figure(figsize=(12, 6))
    df[numeric_columns].boxplot()
    plt.title("数值列的箱线图")
    plt.xticks(rotation=45)
    plt.tight_layout()
    plt.savefig(os.path.join(output_dir, "箱线图.png"))
    plt.close()
    
    # 相关性热图
    plt.figure(figsize=(12, 10))
    sns.heatmap(df[numeric_columns].corr(), annot=True, cmap='coolwarm')
    plt.title("相关性热图")
    plt.tight_layout()
    plt.savefig(os.path.join(output_dir, "相关性热图.png"))
    plt.close()
    
    # 分类列的条形图
    for column in categorical_columns:
        plt.figure(figsize=(12, 6))
        df[column].value_counts().plot(kind='bar')
        plt.title(f"{column}的计数")
        plt.xticks(rotation=45)
        plt.tight_layout()
        plt.savefig(os.path.join(output_dir, f"{column}_条形图.png"))
        plt.close()
    
    # K-means聚类可视化
    if '聚类' in df.columns and 'PCA1' in df.columns and 'PCA2' in df.columns:
        plt.figure(figsize=(10, 8))
        sns.scatterplot(data=df, x='PCA1', y='PCA2', hue='聚类', palette='deep')
        plt.title('K-means聚类结果 (PCA)')
        plt.savefig(os.path.join(output_dir, "K-means聚类.png"))
        plt.close()
    
    # PCA可视化
    if 'PCA1' in df.columns and 'PCA2' in df.columns:
        plt.figure(figsize=(10, 8))
        sns.scatterplot(data=df, x='PCA1', y='PCA2')
        plt.title('PCA: 前两个主成分')
        plt.savefig(os.path.join(output_dir, "PCA可视化.png"))
        plt.close()
    
    # 散点图矩阵
    if len(numeric_columns) >= 2:
        plt.figure(figsize=(15, 15))
        sns.pairplot(df[numeric_columns])
        plt.tight_layout()
        plt.savefig(os.path.join(output_dir, "散点图矩阵.png"))
        plt.close()
    
    # 时间序列图（如果有日期列）
    date_columns = df.select_dtypes(include=['datetime64']).columns
    if len(date_columns) > 0 and len(numeric_columns) > 0:
        date_col = date_columns[0]
        for num_col in numeric_columns[:3]:  # 只绘制前三个数值列
            plt.figure(figsize=(12, 6))
            plt.plot(df[date_col], df[num_col])
            plt.title(f"{num_col}随时间的变化")
            plt.xlabel("日期")
            plt.ylabel(num_col)
            plt.xticks(rotation=45)
            plt.tight_layout()
            plt.savefig(os.path.join(output_dir, f"{num_col}_时间序列图.png"))
            plt.close()
    
    # 饼图（对于分类变量）
    for column in categorical_columns[:3]:  # 只绘制前三个分类列
        plt.figure(figsize=(10, 10))
        df[column].value_counts().plot(kind='pie', autopct='%1.1f%%')
        plt.title(f"{column}的分布")
        plt.ylabel('')
        plt.savefig(os.path.join(output_dir, f"{column}_饼图.png"))
        plt.close()

def export_to_excel(df, analysis, output_file):
    """将处理后的数据和分析结果导出到Excel"""
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='处理后的数据', index=False)
        analysis['基本统计'].to_excel(writer, sheet_name='基本统计')
        analysis['相关性矩阵'].to_excel(writer, sheet_name='相关性矩阵')
        analysis['偏度'].to_excel(writer, sheet_name='偏度', header=False)
        analysis['峰度'].to_excel(writer, sheet_name='峰度', header=False)
        
        for col, counts in analysis['分类变量计数'].items():
            counts.to_excel(writer, sheet_name=f'{col[:31]}计数')
        
        if 'K-means聚类' in analysis:
            analysis['K-means聚类'].to_excel(writer, sheet_name='K-means聚类')
        
        if 'PCA解释方差比' in analysis:
            pd.Series(analysis['PCA解释方差比']).to_excel(writer, sheet_name='PCA解释方差比')
        
        workbook = writer.book
        worksheet = workbook['处理后的数据']
        
        create_excel_charts(worksheet, df)

def create_excel_charts(worksheet, df):
    """根据数据类型创建Excel图表"""
    numeric_columns = df.select_dtypes(include=['int64', 'float64']).columns
    categorical_columns = df.select_dtypes(include=['object']).columns
    date_columns = df.select_dtypes(include=['datetime64']).columns

    for idx, column in enumerate(numeric_columns, start=1):
        # 创建柱状图
        chart = BarChart()
        chart.title = f"{column}分布"
        chart.x_axis.title = column
        chart.y_axis.title = "值"
        data = Reference(worksheet, min_col=idx, min_row=1, max_row=len(df)+1)
        chart.add_data(data, titles_from_data=True)
        worksheet.add_chart(chart, f"{chr(65+idx)}1")

        # 如果有日期列，创建折线图
        if len(date_columns) > 0:
            line_chart = LineChart()
            line_chart.title = f"{column}随时间变化"
            line_chart.x_axis.title = "日期"
            line_chart.y_axis.title = column
            dates = Reference(worksheet, min_col=df.columns.get_loc(date_columns[0])+1, min_row=2, max_row=len(df)+1)
            values = Reference(worksheet, min_col=idx, min_row=1, max_row=len(df)+1)
            line_chart.add_data(values, titles_from_data=True)
            line_chart.set_categories(dates)
            worksheet.add_chart(line_chart, f"{chr(65+idx)}15")

    # 为分类列创建饼图
    for idx, column in enumerate(categorical_columns[:3], start=1):  # 只为前三个分类列创建饼图
        pie_chart = PieChart()
        pie_chart.title = f"{column}分布"
        labels = Reference(worksheet, min_col=df.columns.get_loc(column)+1, min_row=2, max_row=len(df)+1)
        data = Reference(worksheet, min_col=df.columns.get_loc(column)+1, min_row=1, max_row=len(df)+1)
        pie_chart.add_data(data, titles_from_data=True)
        pie_chart.set_categories(labels)
        worksheet.add_chart(pie_chart, f"{chr(65+len(numeric_columns)+idx)}1")

    # 如果有两个或更多数值列，创建散点图
    if len(numeric_columns) >= 2:
        scatter_chart = ScatterChart()
        scatter_chart.title = f"{numeric_columns[0]} vs {numeric_columns[1]}"
        scatter_chart.x_axis.title = numeric_columns[0]
        scatter_chart.y_axis.title = numeric_columns[1]
        xvalues = Reference(worksheet, min_col=df.columns.get_loc(numeric_columns[0])+1, min_row=2, max_row=len(df)+1)
        yvalues = Reference(worksheet, min_col=df.columns.get_loc(numeric_columns[1])+1, min_row=2, max_row=len(df)+1)
        series = Series(yvalues, xvalues, title_from_data=True)
        scatter_chart.series.append(series)
        worksheet.add_chart(scatter_chart, f"{chr(65+len(numeric_columns)+len(categorical_columns)+1)}1")

def process_excel_files(directory, filter_column=None, filter_condition=None, merge_column=None):
    """处理Excel文件的主函数"""
    print("正在加载Excel文件...")
    df = load_excel_files(directory)
    
    print("正在清理数据...")
    df = clean_data(df)
    
    if filter_column and filter_condition:
        print(f"正在筛选数据: {filter_column} {filter_condition}")
        df = filter_data(df, filter_column, filter_condition)
    
    if merge_column:
        print(f"正在基于以下列合并数据: {merge_column}")
        df = merge_data(df, merge_column)
    
    print("正在分析数据...")
    analysis, df = analyze_data(df)
    
    print("正在创建可视化...")
    output_dir = os.path.join(directory, "可视化结果")
    os.makedirs(output_dir, exist_ok=True)
    create_visualizations(df, analysis, output_dir)
    
    print("正在导出结果...")
    output_file = os.path.join(directory, "处理后的主数据.xlsx")
    export_to_excel(df, analysis, output_file)
    
    print(f"处理完成！结果保存在 {output_file} 和 {output_dir}")

def main():
    """主函数"""
    directory = input("请输入包含Excel文件的目录路径: ")
    filter_column = input("请输入要筛选的列名（如果不需要筛选，请直接按回车）: ")
    filter_condition = input("请输入筛选条件（例如: >100, <50, =值, 包含文本）: ") if filter_column else None
    merge_column = input("请输入要合并数据的列名（如果不需要合并，请直接按回车）: ")
    
    process_excel_files(directory, filter_column, filter_condition, merge_column)

if __name__ == "__main__":
    main()

# 使用说明：
# 1. 确保已安装所有必要的库：pip install pandas numpy openpyxl matplotlib seaborn scikit-learn scipy
# 2. 将所有需要处理的Excel文件放在同一个目录下
# 3. 运行脚本，按提示输入相关信息
# 4. 处理完成后，结果将保存在指定目录下的"处理后的主数据.xlsx"文件中，可视化结果将保存在"可视化结果"文件夹中

# 注意事项：
# - 该脚本会处理目录中的所有.xlsx文件
# - 如果数据量很大，处理过程可能需要一些时间
# - 确保有足够的磁盘空间来保存处理后的数据和可视化结果
